"""Dashboard blueprint – analytics, charts and report exports."""
import logging
import io
from collections import defaultdict
from datetime import datetime, timedelta, date
from flask import render_template, request, jsonify, send_file
from flask_login import login_required
from sqlalchemy import func
from ...extensions import db
from ...models.sale import Sale, SaleItem
from ...models.product import Product
from ...models.inventory import InventoryItem
from . import dashboard_bp

logger = logging.getLogger(__name__)


def _date_range(period: str):
    """Return (start_date, end_date) for the requested period."""
    today = date.today()
    if period == "week":
        start = today - timedelta(days=6)
    elif period == "month":
        start = today.replace(day=1)
    else:  # day
        start = today
    return start, today


# ---------------------------------------------------------------------------
# Main dashboard
# ---------------------------------------------------------------------------

@dashboard_bp.route("/")
@login_required
def index():
    today = date.today()
    # KPIs
    sales_today = Sale.query.filter(
        Sale.status == Sale.STATUS_PAID,
        func.date(Sale.created_at) == today,
    ).all()
    total_today = sum(float(s.total) for s in sales_today)
    count_today = len(sales_today)

    # Low stock alerts
    low_stock_items = InventoryItem.query.filter(
        InventoryItem.stock <= InventoryItem.min_stock
    ).count()

    # Expiring soon (next 7 days)
    next_week = today + timedelta(days=7)
    expiring_count = InventoryItem.query.filter(
        InventoryItem.expiry_date.isnot(None),
        InventoryItem.expiry_date <= next_week,
        InventoryItem.expiry_date >= today,
    ).count()

    # Top 5 products today
    top_products = (
        db.session.query(Product.name, func.sum(SaleItem.quantity).label("qty"))
        .join(SaleItem, SaleItem.product_id == Product.id)
        .join(Sale, Sale.id == SaleItem.sale_id)
        .filter(
            Sale.status == Sale.STATUS_PAID,
            func.date(Sale.created_at) == today,
        )
        .group_by(Product.name)
        .order_by(func.sum(SaleItem.quantity).desc())
        .limit(5)
        .all()
    )

    return render_template(
        "dashboard/index.html",
        total_today=total_today,
        count_today=count_today,
        low_stock_items=low_stock_items,
        expiring_count=expiring_count,
        top_products=top_products,
    )


# ---------------------------------------------------------------------------
# JSON endpoints for Chart.js
# ---------------------------------------------------------------------------

@dashboard_bp.route("/api/ventas")
@login_required
def api_sales():
    """Return daily sales totals for the selected period."""
    period = request.args.get("period", "week")
    start, end = _date_range(period)

    rows = (
        db.session.query(
            func.date(Sale.created_at).label("day"),
            func.sum(Sale.total).label("total"),
            func.count(Sale.id).label("count"),
        )
        .filter(
            Sale.status == Sale.STATUS_PAID,
            func.date(Sale.created_at) >= start,
            func.date(Sale.created_at) <= end,
        )
        .group_by(func.date(Sale.created_at))
        .order_by(func.date(Sale.created_at))
        .all()
    )

    data = {
        "labels": [str(r.day) for r in rows],
        "totals": [float(r.total) for r in rows],
        "counts": [int(r.count) for r in rows],
    }
    return jsonify(data)


@dashboard_bp.route("/api/platillos-mas-vendidos")
@login_required
def api_top_products():
    """Top-10 most-sold products in the selected period."""
    period = request.args.get("period", "week")
    start, end = _date_range(period)

    rows = (
        db.session.query(
            Product.name,
            func.sum(SaleItem.quantity).label("qty"),
            func.sum(SaleItem.subtotal).label("revenue"),
        )
        .join(SaleItem, SaleItem.product_id == Product.id)
        .join(Sale, Sale.id == SaleItem.sale_id)
        .filter(
            Sale.status == Sale.STATUS_PAID,
            func.date(Sale.created_at) >= start,
            func.date(Sale.created_at) <= end,
        )
        .group_by(Product.name)
        .order_by(func.sum(SaleItem.quantity).desc())
        .limit(10)
        .all()
    )

    return jsonify(
        {
            "labels": [r.name for r in rows],
            "quantities": [int(r.qty) for r in rows],
            "revenue": [float(r.revenue) for r in rows],
        }
    )


@dashboard_bp.route("/api/rentabilidad")
@login_required
def api_profitability():
    """Revenue vs cost analysis."""
    period = request.args.get("period", "month")
    start, end = _date_range(period)

    rows = (
        db.session.query(
            Product.name,
            func.sum(SaleItem.subtotal).label("revenue"),
            func.sum(SaleItem.quantity * Product.cost).label("cost"),
        )
        .join(SaleItem, SaleItem.product_id == Product.id)
        .join(Sale, Sale.id == SaleItem.sale_id)
        .filter(
            Sale.status == Sale.STATUS_PAID,
            func.date(Sale.created_at) >= start,
            func.date(Sale.created_at) <= end,
        )
        .group_by(Product.name)
        .order_by(func.sum(SaleItem.subtotal).desc())
        .limit(10)
        .all()
    )

    return jsonify(
        {
            "labels": [r.name for r in rows],
            "revenue": [float(r.revenue) for r in rows],
            "cost": [float(r.cost or 0) for r in rows],
        }
    )


# ---------------------------------------------------------------------------
# Full analytics view
# ---------------------------------------------------------------------------

@dashboard_bp.route("/analitica")
@login_required
def analytics():
    return render_template("dashboard/analytics.html")


# ---------------------------------------------------------------------------
# Report exports
# ---------------------------------------------------------------------------

@dashboard_bp.route("/reporte/pdf")
@login_required
def report_pdf():
    """Export a sales summary PDF for the selected period."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        return "reportlab no está instalado.", 500

    period = request.args.get("period", "week")
    start, end = _date_range(period)

    sales = Sale.query.filter(
        Sale.status == Sale.STATUS_PAID,
        func.date(Sale.created_at) >= start,
        func.date(Sale.created_at) <= end,
    ).order_by(Sale.created_at.desc()).all()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("FondaControl – Reporte de Ventas", styles["Title"]))
    elements.append(Paragraph(f"Período: {start} al {end}", styles["Normal"]))
    elements.append(Spacer(1, 12))

    table_data = [["Folio", "Fecha", "Total", "Método", "Cajero"]]
    for s in sales:
        table_data.append([
            s.folio,
            s.created_at.strftime("%d/%m/%Y %H:%M"),
            f"${float(s.total):.2f}",
            s.payment_method,
            s.cashier.username if s.cashier else "—",
        ])

    table = Table(table_data)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e67e22")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fef9f5")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    elements.append(table)

    total_amount = sum(float(s.total) for s in sales)
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"<b>Total: ${total_amount:.2f}</b>", styles["Normal"]))

    doc.build(elements)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"reporte_ventas_{start}_{end}.pdf",
    )


@dashboard_bp.route("/reporte/excel")
@login_required
def report_excel():
    """Export sales to Excel."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return "openpyxl no está instalado.", 500

    period = request.args.get("period", "week")
    start, end = _date_range(period)

    sales = Sale.query.filter(
        Sale.status == Sale.STATUS_PAID,
        func.date(Sale.created_at) >= start,
        func.date(Sale.created_at) <= end,
    ).order_by(Sale.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas"

    header_fill = PatternFill("solid", fgColor="E67E22")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["Folio", "Fecha", "Total", "Subtotal", "Método de pago", "Mesa", "Cajero"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row, s in enumerate(sales, 2):
        ws.append([
            s.folio,
            s.created_at.strftime("%d/%m/%Y %H:%M"),
            float(s.total),
            float(s.subtotal),
            s.payment_method,
            s.table_number or "—",
            s.cashier.username if s.cashier else "—",
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"reporte_ventas_{start}_{end}.xlsx",
    )
